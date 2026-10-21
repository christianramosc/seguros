# -*- coding: utf-8 -*-
"""
Cotizaciones de póliza externa (Multicotizador de Istra) para comparar contra la póliza actual de cada crédito.

El archivo es un Excel con una hoja «Comparativo» (o «Cotizaciones» dentro de Polizas_extraidas.xlsx): fila de
encabezados con CLIENTE y SERIE, una columna por aseguradora terminada en «(VC)» (valor comercial) y, si hay,
en «(VF)» (valor factura), más NOTAS. La más barata se calcula aquí, no se toma de las fórmulas del archivo.

  leer(xlsx_bytes)            → DataFrame por SERIE con precios, la más barata y notas (o None)
  agregar(base, cot)          → base + EXTERNA MÁS BARATA, ASEGURADORA EXTERNA, AHORRO EXTERNA, …
  incrustar(excel, cot_bytes) → Polizas_extraidas.xlsx con la hoja «Cotizaciones» (así se conserva)
"""
import io
import re

import pandas as pd

HOJA = "Cotizaciones"
HOJAS_FUENTE = ("Comparativo", HOJA)
# Aseguradoras que se muestran en el PDF; las tarifas ABA (Chubb) se juntan en una sola (la más barata).
GRUPOS = {"Quálitas": "QUÁLITAS", "ABA": "ABA/CHUBB", "AXA": "AXA", "GNP": "GNP", "HDI": "HDI", "Zurich": "ZURICH"}
COLS = ["EXTERNA MÁS BARATA", "ASEGURADORA EXTERNA", "AHORRO EXTERNA", "EXTERNA VF MÁS BARATA", "NOTA COTIZACIÓN",
        "FECHA COTIZACIÓN"]


def _hoja(xl):
    return next((h for h in HOJAS_FUENTE if h in xl.sheet_names), None)


def _num(v):
    try:
        f = float(v)
        return f if f == f and f > 0 else None
    except (TypeError, ValueError):
        return None


def _grupo(col):
    nombre = col.rsplit("(", 1)[0].strip()
    return next((g for k, g in GRUPOS.items() if nombre.upper().startswith(k.upper())), nombre.upper())


def leer(xlsx_bytes):
    """DataFrame indexado por SERIE, o None si el archivo no trae cotizaciones."""
    if not xlsx_bytes:
        return None
    xl = pd.ExcelFile(io.BytesIO(xlsx_bytes))
    hoja = _hoja(xl)
    if not hoja:
        return None
    crudo = pd.read_excel(xl, hoja, header=None, dtype=object)
    fila = next((i for i, r in crudo.iterrows()
                 if {"CLIENTE", "SERIE"} <= {str(v).strip().upper() for v in r.values}), None)
    if fila is None:
        return None
    fecha = ""
    for _, r in crudo.iloc[:fila].iterrows():
        vals = [str(v).strip() for v in r.values if str(v) not in ("nan", "None")]
        if vals and vals[0].upper().startswith("FECHA COTIZ") and len(vals) > 1:
            fecha = vals[1][:10]
    df = crudo.iloc[fila + 1:].copy()
    df.columns = [str(c).strip() for c in crudo.iloc[fila]]
    df = df[df["SERIE"].notna() & df["SERIE"].astype(str).str.match(r"^[A-Z0-9]{17}$")]
    vc = [c for c in df.columns if c.endswith("(VC)") and not re.match(r"(MÁS BARATA|ASEGURADORA)", c)]
    vf = [c for c in df.columns if c.endswith("(VF)") and not c.startswith("MÁS BARATA")]
    out = []
    for _, r in df.iterrows():
        precios = {}
        for c in vc:                                        # por aseguradora, la tarifa más barata
            v = _num(r[c])
            g = _grupo(c)
            if v and (g not in precios or v < precios[g]):
                precios[g] = v
        en_vf = [v for v in (_num(r[c]) for c in vf) if v]
        vf_min = min(en_vf) if en_vf else None
        mejor = min(precios, key=precios.get) if precios else ""
        nota = str(r.get("NOTAS") or "").strip()
        out.append({"SERIE": str(r["SERIE"]).strip(), "PRECIOS": precios,
                    "EXTERNA MÁS BARATA": precios.get(mejor), "ASEGURADORA EXTERNA": mejor,
                    "EXTERNA VF MÁS BARATA": vf_min, "NOTA COTIZACIÓN": "" if nota in ("nan", "None") else nota,
                    "FECHA COTIZACIÓN": fecha})
    return pd.DataFrame(out).drop_duplicates("SERIE").set_index("SERIE") if out else None


def agregar(base, cot):
    """Columnas de la cotización en la base (por SERIE). AHORRO = prima anual actual − la externa más barata."""
    b = base.copy()
    if cot is None or cot.empty:
        for c in COLS + ["PRECIOS EXTERNA"]:
            b[c] = None
        return b
    for c in COLS:
        if c in cot.columns:
            b[c] = b["SERIE"].map(cot[c])
    b["PRECIOS EXTERNA"] = b["SERIE"].map(cot["PRECIOS"])
    b["AHORRO EXTERNA"] = [round(a - e, 2) if a == a and a and e == e and e else None
                           for a, e in zip(b["PRIMA ANUAL ACTUAL"], b["EXTERNA MÁS BARATA"])]
    return b


def incrustar(excel_bytes, cot_bytes):
    """Copia la hoja de cotizaciones (solo valores) dentro del Excel de pólizas como «Cotizaciones»."""
    from openpyxl import load_workbook
    origen = load_workbook(io.BytesIO(cot_bytes), data_only=True)
    ws_o = origen[next(h for h in HOJAS_FUENTE if h in origen.sheetnames)]
    wb = load_workbook(io.BytesIO(excel_bytes))
    if HOJA in wb.sheetnames:
        del wb[HOJA]
    copiar_hoja(ws_o, wb.create_sheet(HOJA))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def copiar_hoja(ws_o, ws_d):
    """Valores, anchos y formato básico de una hoja a otra (sirve entre libros)."""
    from copy import copy
    for fila in ws_o.iter_rows():
        for c in fila:
            d = ws_d.cell(c.row, c.column, c.value)
            if c.has_style:
                d.font, d.fill, d.number_format = copy(c.font), copy(c.fill), c.number_format
                d.alignment, d.border = copy(c.alignment), copy(c.border)
    for k, dim in ws_o.column_dimensions.items():
        ws_d.column_dimensions[k].width = dim.width
    ws_d.freeze_panes = ws_o.freeze_panes
